"""把商业计划书从 Markdown 转成可提交的文档。

产出两份：
  * .html —— 带打印样式，浏览器里 Ctrl+P 直接存成 PDF，排版可控
  * .docx —— 可编辑，用来补团队信息之类的占位内容

用法：
    python scripts/make_bp_doc.py ../outputs/靶点尽调_商业计划书.md
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

PRINT_CSS = """
@page { size: A4; margin: 20mm 18mm; }
* { box-sizing: border-box; }
body {
  font-family: "PingFang SC", "Microsoft YaHei", "Source Han Sans SC", -apple-system, sans-serif;
  color: #12161c; line-height: 1.75; font-size: 10.5pt;
  max-width: 180mm; margin: 0 auto; padding: 8mm 0;
}
h1 { font-size: 20pt; margin: 0 0 6pt; letter-spacing: -.01em; }
h2 { font-size: 14pt; margin: 22pt 0 8pt; padding-bottom: 5pt; border-bottom: 1.5pt solid #1f4fd1; color: #1f4fd1; page-break-after: avoid; }
h3 { font-size: 11.5pt; margin: 14pt 0 6pt; page-break-after: avoid; }
p { margin: 0 0 8pt; }
strong { font-weight: 600; }
ol, ul { margin: 0 0 9pt; padding-left: 18pt; }
li { margin: 3pt 0; }
table { border-collapse: collapse; width: 100%; margin: 8pt 0 12pt; font-size: 9.5pt; page-break-inside: avoid; }
th, td { border: .6pt solid #d7dee8; padding: 5pt 8pt; text-align: left; vertical-align: top; }
th { background: #f2f5fa; font-weight: 600; }
code { background: #f2f5fa; padding: 1pt 4pt; border-radius: 3pt; font-size: 9.5pt; }
a { color: #1f4fd1; text-decoration: none; word-break: break-all; }
blockquote { margin: 8pt 0; padding: 6pt 12pt; border-left: 3pt solid #1f4fd1; background: #f7f9fd; color: #3a4452; }
hr { border: 0; border-top: .6pt solid #d7dee8; margin: 16pt 0; }
@media print { a { color: #12161c; } }
"""

HTML_TEMPLATE = """<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>{css}</style>
</head>
<body>
{body}
</body>
</html>
"""


def markdown_to_html(text: str, title: str) -> str:
    try:
        from markdown_it import MarkdownIt
    except ImportError:
        print("需要 markdown_it：pip install markdown-it-py", file=sys.stderr)
        raise
    md = MarkdownIt("commonmark", {"html": False}).enable("table")
    return HTML_TEMPLATE.format(title=title, css=PRINT_CSS, body=md.render(text))


def _set_cjk_font(run, name: str = "微软雅黑") -> None:
    """python-docx 对中日韩字体需要单独设置 eastAsia，否则会回退成宋体。"""
    from docx.oxml.ns import qn

    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def _add_runs(paragraph, text: str) -> None:
    """处理 **加粗** 与 `代码` 两种行内标记。"""
    for piece in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
        else:
            run = paragraph.add_run(piece)
        _set_cjk_font(run)


def markdown_to_docx(text: str, out_path: Path, title: str) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(6)

    lines = text.split("\n")
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        # 表格：连续的 | 开头行
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [
                [c.strip() for c in r.strip("|").split("|")]
                for r in block
                if not re.match(r"^\|[\s:\-|]+\|$", r)
            ]
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row[: len(rows[0])]):
                        paragraph = table.cell(r, c).paragraphs[0]
                        _add_runs(paragraph, cell)
                        for run in paragraph.runs:
                            run.font.size = Pt(9.5)
                            if r == 0:
                                run.bold = True
            continue

        if not line.strip():
            i += 1
            continue
        if re.match(r"^---+$", line.strip()):
            i += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            paragraph = doc.add_heading(level=min(level, 4))
            _add_runs(paragraph, heading.group(2))
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0x1F, 0x4F, 0xD1) if level == 2 else RGBColor(0x12, 0x16, 0x1C)
            i += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", line)
        numbered = re.match(r"^(\d+)[.)]\s+(.*)$", line)
        if bullet:
            _add_runs(doc.add_paragraph(style="List Bullet"), bullet.group(1))
            i += 1
            continue
        if numbered:
            _add_runs(doc.add_paragraph(style="List Number"), numbered.group(2))
            i += 1
            continue

        if line.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(18)
            _add_runs(paragraph, line[2:])
            for run in paragraph.runs:
                run.italic = True
            i += 1
            continue

        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _add_runs(paragraph, line)
        i += 1

    doc.save(str(out_path))


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="把商业计划书转成 HTML 与 Word")
    parser.add_argument("source", help="Markdown 源文件")
    parser.add_argument("--out-dir", help="输出目录（默认与源文件同目录）")
    args = parser.parse_args(argv)

    source = Path(args.source)
    text = source.read_text(encoding="utf-8")
    title = next((l.lstrip("# ").strip() for l in text.split("\n") if l.startswith("# ")), source.stem)
    out_dir = Path(args.out_dir) if args.out_dir else source.parent
    out_dir.mkdir(parents=True, exist_ok=True)

    html_path = out_dir / (source.stem + ".html")
    html_path.write_text(markdown_to_html(text, title), encoding="utf-8")
    print(f"  HTML  {html_path}  （浏览器打开后 Ctrl+P 可存为 PDF）")

    docx_path = out_dir / (source.stem + ".docx")
    markdown_to_docx(text, docx_path, title)
    print(f"  DOCX  {docx_path}  （可编辑）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
